"""Create V3 from the user-edited V2 while preserving its formatting."""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

import pymupdf
from docx import Document


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Chapter_Roe_RAR_V2.docx"
TARGET = HERE / "Chapter_Roe_RAR_V3.docx"
PDF = TARGET.with_suffix(".pdf")

FIGURES_WITH_DESIGN = [
    "wave_system.png",
    "experimental_design.png",
    "anchor_patterns.png",
    "core_profiles.png",
    "error_evolution.png",
    "seed_results.png",
    "roe_profiles.png",
    "hydraulic_diagnostics.png",
]

FIGURES_WITHOUT_DESIGN = [
    "wave_system.png",
    "anchor_patterns.png",
    "core_profiles.png",
    "error_evolution.png",
    "seed_results.png",
    "roe_profiles.png",
    "hydraulic_diagnostics.png",
]


ABSTRACT_OPENING = (
    "Dam-break floods generate rapidly propagating waves that can threaten "
    "downstream communities, infrastructure, and hydraulic structures. "
    "Predicting their arrival time, depth, and velocity requires a model that "
    "can represent a moving shock wave while respecting conservation across "
    "the front. This chapter examines one question: under a fixed architecture "
    "and training budget, how do the spatial residual and the placement of "
    "collocation points affect a neural solution of the Stoker wet-bed "
    "dam-break problem? Six configurations cross a pointwise differential "
    "residual or a Roe numerical-flux residual with uniform enrichment, "
    "residual-based adaptive refinement (RAR), and a hybrid sampling rule. "
    "Each configuration is trained from three matched random seeds. All "
    "quantitative results come from the 18 runs contained in runs.tar.gz."
)

INTRODUCTION_OPENING = (
    "Dam-break floods are among the most rapid and destructive unsteady flows "
    "in hydraulic engineering. A sudden release of impounded water can "
    "produce large depths, high velocities, and short warning times "
    "downstream, placing people, bridges, roads, buildings, and other critical "
    "infrastructure at risk. Reliable prediction of flood-wave arrival time "
    "and intensity is therefore important for hazard mapping, emergency "
    "planning, dam-safety assessment, and the design of protective measures "
    "[1-4]."
)

INTRODUCTION_PROBLEM = (
    "The hydraulic response is difficult to calculate because a dam break "
    "produces several wave structures at once. In the idealized wet-bed "
    "problem, sudden removal of the barrier creates a left-going rarefaction "
    "wave and a right-going shock wave, separated by a nearly uniform region "
    "of moving water. The shock position determines when the rapid flow change "
    "reaches a downstream location, while the depth and velocity behind the "
    "shock determine the transported discharge and momentum. The wet-bed dam "
    "break is consequently a compact but exacting benchmark: its rarefaction "
    "is smooth, whereas the shock wave is a discontinuity whose speed and jump "
    "are constrained by conservation. A low average depth error can still "
    "conceal an incorrect shock location or velocity because the upstream "
    "reservoir dominates the depth norm [1-4]. This chapter therefore focuses "
    "on the governing-equation residual and the distribution of its evaluation "
    "points."
)

REPLACEMENTS = [
    ("hydraulic-bore", "shock-wave"),
    ("Hydraulic-bore", "Shock-wave"),
    ("hydraulic bore", "shock wave"),
    ("Hydraulic bore", "Shock wave"),
    ("bore-location", "shock-location"),
    ("Bore-location", "Shock-location"),
    ("bore detection", "shock-wave detection"),
    ("Bore detection", "Shock-wave detection"),
    ("bore speed", "shock-wave speed"),
    ("Bore speed", "Shock-wave speed"),
    ("bore position", "shock position"),
    ("Bore position", "Shock position"),
    ("bore error", "shock-location error"),
    ("Bore error", "Shock-location error"),
    ("bore crossings", "shock crossings"),
    ("Bore crossings", "Shock crossings"),
    ("bore crossing", "shock crossing"),
    ("Bore crossing", "Shock crossing"),
    ("Valid bore", "Valid shock"),
    ("valid bore", "valid shock"),
    ("bore", "shock wave"),
    ("Bore", "Shock wave"),
]


def iter_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph


def replace_terms(document):
    for paragraph in iter_paragraphs(document):
        for run in paragraph.runs:
            text = run.text
            for old, new in REPLACEMENTS:
                text = text.replace(old, new)
            if text != run.text:
                run.text = text


def replace_figures(document):
    figure_sets = {
        len(FIGURES_WITH_DESIGN): FIGURES_WITH_DESIGN,
        len(FIGURES_WITHOUT_DESIGN): FIGURES_WITHOUT_DESIGN,
    }
    figures = figure_sets.get(len(document.inline_shapes))
    if figures is None:
        raise RuntimeError(
            f"Expected seven or eight figures, found {len(document.inline_shapes)}"
        )
    for shape, filename in zip(document.inline_shapes, figures):
        blip = shape._inline.graphic.graphicData.pic.blipFill.blip
        relationship_id = blip.embed
        image_part = document.part.related_parts[relationship_id]
        image_part._blob = (HERE / "focused_figures" / filename).read_bytes()


def find_heading(document, heading):
    for index, paragraph in enumerate(document.paragraphs):
        if paragraph.text.strip() == heading:
            return index
    raise ValueError(f"Heading not found: {heading}")


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)

    subprocess.run(
        ["python", str(HERE / "analyze_focused_archive.py")],
        cwd=HERE.parent.parent,
        check=True,
    )

    document = Document(SOURCE)
    abstract = find_heading(document, "Abstract")
    introduction = find_heading(document, "1. Introduction")
    document.paragraphs[abstract + 1].text = ABSTRACT_OPENING
    document.paragraphs[introduction + 1].text = INTRODUCTION_OPENING
    document.paragraphs[introduction + 2].text = INTRODUCTION_PROBLEM

    replace_terms(document)
    replace_figures(document)
    document.core_properties.subject = (
        "Dam-break shock-wave modeling with Roe-flux physics-informed neural networks"
    )
    document.save(TARGET)

    subprocess.run(
        ["python", str(HERE / "render_chapter.py"), str(TARGET)],
        cwd=HERE.parent.parent,
        check=True,
    )

    rendered = pymupdf.open(PDF)
    full_text = "\n".join(page.get_text() for page in rendered)
    if re.search(r"\bbore\b", full_text, flags=re.IGNORECASE):
        raise AssertionError("The rendered document still contains 'bore'.")
    intro_page = next(
        index + 1 for index, page in enumerate(rendered)
        if "1. Introduction" in page.get_text()
    )
    method_page = next(
        index + 1 for index, page in enumerate(rendered)
        if "2. Method" in page.get_text()
    )
    if not 20 <= len(rendered) <= 25:
        raise AssertionError(f"V3 has {len(rendered)} pages; expected 20-25.")
    if method_page - intro_page + 1 > 5:
        raise AssertionError("The introduction exceeds five pages.")
    print(
        f"Created {TARGET}\n"
        f"Pages: {len(rendered)}\n"
        f"Introduction: pages {intro_page}-{method_page} (Method starts on {method_page})\n"
        "Terminology check: no instances of 'bore'"
    )


if __name__ == "__main__":
    main()
