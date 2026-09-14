"""Replace the dense shock infographic with three simple, separate figures."""
from __future__ import annotations

from copy import deepcopy
import json
from pathlib import Path

import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Chapter_Roe_RAR_V3_visualized.docx"
OUTPUT = HERE / "Chapter_Roe_RAR_V3_clear_figures.docx"
FIGURE_DIR = HERE / "focused_figures"
FIGURE_A = FIGURE_DIR / "shock_midpoint_clear.png"
FIGURE_B = FIGURE_DIR / "shock_crossing_clear.png"
FIGURE_C = FIGURE_DIR / "shock_position_error_clear.png"

BLUE = "#0072B2"
ORANGE = "#D55E00"
PURPLE = "#7B3294"
RED = "#B2182B"
GRAY = "#52606D"
LIGHT_GRID = "#D7DCE2"


def load_example():
    run_dir = HERE.parent.parent / "runs_extracted" / "runs" / "ablation" / "roe_roe_rar__seed0"
    with np.load(run_dir / "prediction_t1.2.npz") as data:
        x = np.asarray(data["x"])
        h = np.asarray(data["h"])
    result = json.loads((run_dir / "result.json").read_text(encoding="utf-8"))
    h_star = 0.3961748167990253
    h_right = 0.1
    midpoint = (h_star + h_right) / 2
    crossings = np.where(
        (x[:-1] > 0.5) & (h[:-1] >= midpoint) & (h[1:] < midpoint)
    )[0]
    if len(crossings) != 1:
        raise RuntimeError(f"Expected one descending crossing; found {len(crossings)}")
    i = int(crossings[0])
    return {
        "x": x,
        "h": h,
        "i": i,
        "h_star": h_star,
        "h_right": h_right,
        "midpoint": midpoint,
        "predicted": float(result["shock_prediction_m_t1.2"]),
        "exact": float(result["shock_exact_m_t1.2"]),
        "error": float(result["shock_location_error_m_t1.2"]),
    }


def base_style(axis) -> None:
    axis.grid(True, color=LIGHT_GRID, linewidth=0.8, alpha=0.9)
    axis.spines[["top", "right"]].set_visible(False)
    axis.tick_params(labelsize=11)


def figure_midpoint(example) -> None:
    fig, ax = plt.subplots(figsize=(7.0, 3.0))
    upper = example["h_star"]
    lower = example["h_right"]
    middle = example["midpoint"]

    ax.set_title("(a) Find the depth halfway between the two water levels",
                 fontsize=15, fontweight="bold", pad=12)
    ax.set_xlim(0, 10)
    ax.set_ylim(0.06, 0.44)
    ax.set_xticks([])
    ax.set_ylabel("Water depth (m)", fontsize=12)
    ax.spines[["top", "right", "bottom"]].set_visible(False)
    ax.tick_params(axis="y", labelsize=11)

    ax.hlines(upper, 1.4, 8.8, color=BLUE, linewidth=3)
    ax.hlines(lower, 1.4, 8.8, color=BLUE, linewidth=3)
    ax.hlines(middle, 1.4, 8.8, color=ORANGE, linewidth=3, linestyle="--")
    ax.text(8.95, upper, "upper depth = 0.396 m", va="center", fontsize=12, color=BLUE)
    ax.text(8.95, lower, "downstream depth = 0.100 m", va="center", fontsize=12, color=BLUE)
    ax.text(8.95, middle, "midpoint = 0.248 m", va="center", fontsize=12,
            color=ORANGE, fontweight="bold")

    ax.annotate("", xy=(2.3, upper), xytext=(2.3, middle),
                arrowprops=dict(arrowstyle="<->", color=GRAY, linewidth=2))
    ax.annotate("", xy=(2.3, middle), xytext=(2.3, lower),
                arrowprops=dict(arrowstyle="<->", color=GRAY, linewidth=2))
    ax.text(2.55, (upper + middle) / 2, "same distance", va="center", fontsize=11, color=GRAY)
    ax.text(2.55, (middle + lower) / 2, "same distance", va="center", fontsize=11, color=GRAY)
    fig.subplots_adjust(left=0.11, right=0.74, top=0.80, bottom=0.12)
    fig.savefig(FIGURE_A, dpi=240, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def figure_crossing(example) -> None:
    x = example["x"]
    h = example["h"]
    i = example["i"]
    midpoint = example["midpoint"]
    predicted = example["predicted"]
    left_x, right_x = x[i], x[i + 1]
    left_h, right_h = h[i], h[i + 1]

    fig, ax = plt.subplots(figsize=(7.0, 3.8))
    ax.set_title("(b) Locate where the predicted profile crosses the midpoint",
                 fontsize=15, fontweight="bold", pad=12)
    neighborhood = slice(i - 2, i + 4)
    ax.plot(x[neighborhood], h[neighborhood], color=BLUE, linewidth=3)
    ax.plot([left_x, right_x], [left_h, right_h], color=PURPLE, linewidth=4)
    ax.scatter([left_x, right_x], [left_h, right_h], s=90, color=BLUE,
               edgecolor="white", linewidth=1.2, zorder=5)
    ax.axhline(midpoint, color=ORANGE, linestyle="--", linewidth=2.5)
    ax.axvline(predicted, color=PURPLE, linestyle=":", linewidth=2.5)
    ax.scatter([predicted], [midpoint], marker="*", s=230, color=ORANGE,
               edgecolor="white", linewidth=1, zorder=6)

    ax.annotate("above midpoint", xy=(left_x, left_h),
                xytext=(left_x - 0.0016, left_h + 0.010), fontsize=11, ha="center",
                arrowprops=dict(arrowstyle="->", color=GRAY, linewidth=1.5))
    ax.annotate("below midpoint", xy=(right_x, right_h),
                xytext=(right_x + 0.0012, right_h - 0.011), fontsize=11, ha="center",
                arrowprops=dict(arrowstyle="->", color=GRAY, linewidth=1.5))
    ax.annotate("predicted shock\n3.7548 m", xy=(predicted, midpoint),
                xytext=(predicted - 0.0025, midpoint + 0.011), fontsize=12,
                color=PURPLE, fontweight="bold", ha="center",
                arrowprops=dict(arrowstyle="->", color=PURPLE, linewidth=1.8))
    ax.text(left_x - 0.0037, midpoint + 0.001, "midpoint depth",
            color=ORANGE, fontsize=11, va="bottom", fontweight="bold")

    ax.set_xlim(left_x - 0.004, right_x + 0.004)
    ax.set_ylim(right_h - 0.018, left_h + 0.022)
    ax.set_xlabel("Position along the channel (m)", fontsize=12)
    ax.set_ylabel("Predicted water depth (m)", fontsize=12)
    base_style(ax)
    fig.subplots_adjust(left=0.13, right=0.97, top=0.82, bottom=0.18)
    fig.savefig(FIGURE_B, dpi=240, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def figure_position_error(example) -> None:
    exact = example["exact"]
    predicted = example["predicted"]
    error_cm = 100 * example["error"]

    fig, ax = plt.subplots(figsize=(7.0, 2.8))
    ax.set_title("(c) Position error is the gap between the two shock locations",
                 fontsize=15, fontweight="bold", pad=12)
    ax.set_xlim(3.715, 3.765)
    ax.set_ylim(0, 1)
    ax.set_yticks([])
    ax.set_xlabel("Position along the channel (m)", fontsize=12)
    ax.tick_params(axis="x", labelsize=11)
    ax.spines[["top", "right", "left"]].set_visible(False)

    ax.axvline(exact, ymin=0.12, ymax=0.86, color="#222222", linestyle=":", linewidth=3)
    ax.axvline(predicted, ymin=0.12, ymax=0.86, color=PURPLE, linestyle="-.", linewidth=3)
    ax.text(exact, 0.90, "exact shock\n3.7262 m", ha="center", fontsize=12)
    ax.text(predicted, 0.90, "predicted shock\n3.7548 m", ha="center",
            fontsize=12, color=PURPLE)
    ax.annotate("", xy=(exact, 0.48), xytext=(predicted, 0.48),
                arrowprops=dict(arrowstyle="<->", color=RED, linewidth=3))
    ax.text((exact + predicted) / 2, 0.56, f"position error = {error_cm:.2f} cm",
            ha="center", fontsize=14, color=RED, fontweight="bold")
    ax.text((exact + predicted) / 2, 0.22,
            "This distance tells us how far the predicted shock is from where it should be.",
            ha="center", fontsize=11, color=GRAY)

    fig.subplots_adjust(left=0.08, right=0.97, top=0.78, bottom=0.22)
    fig.savefig(FIGURE_C, dpi=240, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def make_figures() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    example = load_example()
    figure_midpoint(example)
    figure_crossing(example)
    figure_position_error(example)


def add_math(paragraph, element) -> None:
    paragraph._p.append(deepcopy(element))


def simplify_explanation(paragraph) -> None:
    equations = paragraph._p.xpath("./m:oMath")
    if not equations:
        raise RuntimeError("The evaluation paragraph does not contain its E equation")
    paragraph.clear()
    paragraph.add_run("Final depth and velocity errors are reported separately because ")
    add_math(paragraph, equations[0])
    paragraph.add_run(
        " can hide which field controls a difference. Figures 3(a)-3(c) show the shock "
        "measurement visually at t = 1.2 s. First, the orange threshold is placed halfway "
        "between the depth behind the shock and the downstream depth. Next, the predicted "
        "shock is the point where the predicted depth profile falls through that orange "
        "line. The two nearest grid samples are joined by a straight segment so that the "
        "crossing can be located between them. Finally, the position error is simply the "
        "horizontal gap between the predicted shock and the exact shock. In the displayed "
        "Roe-R-RAR seed 0 run, that gap is 2.87 cm. The result is accepted only when the "
        "profile has one downward crossing; no crossing or several crossings is invalid."
    )


def add_figure_before(paragraph, image_path: Path, caption_text: str) -> None:
    figure_paragraph = paragraph.insert_paragraph_before(style="Captioned Figure")
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure_paragraph.paragraph_format.keep_with_next = True
    figure_paragraph.add_run().add_picture(str(image_path), width=Inches(5.95))
    caption = paragraph.insert_paragraph_before(style="Image Caption")
    caption.paragraph_format.keep_together = True
    caption.paragraph_format.keep_with_next = False
    caption.add_run(caption_text)


def revise_document() -> None:
    document = Document(SOURCE)
    target = next(
        (p for p in document.paragraphs if "Figure 3 shows the shock-location calculation" in p.text),
        None,
    )
    if target is None:
        raise RuntimeError("Could not find the current shock-detection explanation")
    simplify_explanation(target)

    old_figure = target._p.getnext()
    if old_figure is None or not old_figure.xpath(".//w:drawing"):
        raise RuntimeError("Could not find the old Figure 3 image")
    old_caption = old_figure.getnext()
    if old_caption is None:
        raise RuntimeError("Could not find the old Figure 3 caption")
    following = old_caption.getnext()
    parent = old_figure.getparent()
    parent.remove(old_figure)
    parent.remove(old_caption)
    following_paragraph = next(p for p in document.paragraphs if p._p is following)

    add_figure_before(
        following_paragraph,
        FIGURE_A,
        "Figure 3(a). The orange dashed line is halfway between the upper and downstream "
        "water depths. This midpoint becomes the depth level used to search for the shock.",
    )
    add_figure_before(
        following_paragraph,
        FIGURE_B,
        "Figure 3(b). The predicted shock is where the descending depth profile crosses "
        "the orange midpoint line. The star marks the crossing between the two neighboring "
        "grid samples.",
    )
    add_figure_before(
        following_paragraph,
        FIGURE_C,
        "Figure 3(c). Position error is the horizontal distance between the predicted and "
        "exact shock locations. The displayed run places the shock 2.87 cm too far downstream.",
    )

    document.save(OUTPUT)


def main() -> None:
    make_figures()
    revise_document()
    print(f"Created {OUTPUT}")
    print(f"Created {FIGURE_A}")
    print(f"Created {FIGURE_B}")
    print(f"Created {FIGURE_C}")


if __name__ == "__main__":
    main()
