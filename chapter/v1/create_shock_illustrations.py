"""Replace calculation plots with three distinct conceptual illustrations."""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle
import numpy as np
from PIL import Image
from docx import Document
from docx.shared import Inches


HERE = Path(__file__).resolve().parent
SOURCE = HERE / "Chapter_Roe_RAR_V3_clear_figures.docx"
OUTPUT = HERE / "Chapter_Roe_RAR_V3_new_illustrations.docx"
FIGURE_DIR = HERE / "focused_figures"
FIGURE_A = FIGURE_DIR / "shock_midpoint_illustration.png"
FIGURE_B = FIGURE_DIR / "shock_crossing_illustration.png"
FIGURE_C = FIGURE_DIR / "shock_error_illustration.png"

BLUE = "#0072B2"
WATER = "#A9D8F5"
DEEP_WATER = "#6DB7E5"
ORANGE = "#D55E00"
PURPLE = "#7B3294"
RED = "#B2182B"
DARK = "#263442"
GRAY = "#637282"
PALE_GRAY = "#EDF1F4"


def draw_tank(axis, x: float, water_fraction: float, title: str, value: str,
              *, highlight: bool = False) -> None:
    bottom, width, height = 0.75, 2.15, 3.0
    edge = ORANGE if highlight else DARK
    fill = "#FFD9B8" if highlight else WATER
    tank = FancyBboxPatch(
        (x, bottom), width, height,
        boxstyle="round,pad=0.02,rounding_size=0.10",
        facecolor="white", edgecolor=edge, linewidth=2.6,
    )
    axis.add_patch(tank)
    water_height = height * water_fraction
    water = Rectangle((x + 0.06, bottom + 0.06), width - 0.12, water_height - 0.06,
                      facecolor=fill, edgecolor="none")
    water.set_clip_path(tank)
    axis.add_patch(water)
    surface_y = bottom + water_height
    wave_x = np.linspace(x + 0.08, x + width - 0.08, 80)
    wave_y = surface_y + 0.035 * np.sin(np.linspace(0, 5 * np.pi, 80))
    axis.plot(wave_x, wave_y, color=edge if highlight else BLUE, linewidth=2.5)
    axis.text(x + width / 2, 0.43, title, ha="center", va="center",
              fontsize=12, fontweight="bold", color=edge)
    if water_fraction > 0.75:
        value_y, value_va, value_color = surface_y - 0.34, "top", DARK
    else:
        value_y, value_va, value_color = surface_y + 0.28, "bottom", edge
    axis.text(x + width / 2, value_y, value, ha="center", va=value_va,
              fontsize=13, fontweight="bold", color=value_color)


def make_midpoint_illustration() -> None:
    fig, ax = plt.subplots(figsize=(7.2, 3.3))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis("off")
    ax.set_title("(a) Use the two water depths to choose a halfway level",
                 fontsize=15, fontweight="bold", pad=10)

    maximum = 0.45
    draw_tank(ax, 0.55, 0.396175 / maximum, "behind the shock", "0.396 m")
    draw_tank(ax, 4.15, 0.100000 / maximum, "downstream", "0.100 m")
    draw_tank(ax, 8.85, 0.248088 / maximum, "halfway level", "0.248 m", highlight=True)

    ax.add_patch(FancyArrowPatch((2.85, 2.95), (8.65, 2.55),
                                 connectionstyle="arc3,rad=-0.18", arrowstyle="-|>",
                                 mutation_scale=18, linewidth=2.2, color=GRAY))
    ax.add_patch(FancyArrowPatch((6.40, 1.25), (8.65, 2.15),
                                 connectionstyle="arc3,rad=0.18", arrowstyle="-|>",
                                 mutation_scale=18, linewidth=2.2, color=GRAY))
    ax.text(7.35, 3.55, "pick the level exactly\nhalfway between them",
            ha="center", va="center", fontsize=12, color=DARK,
            bbox=dict(boxstyle="round,pad=0.35", facecolor=PALE_GRAY, edgecolor="none"))
    fig.subplots_adjust(left=0.025, right=0.98, top=0.84, bottom=0.04)
    fig.savefig(FIGURE_A, dpi=240, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def make_crossing_illustration() -> None:
    fig, ax = plt.subplots(figsize=(7.2, 3.5))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis("off")
    ax.set_title("(b) Follow the predicted water surface until it meets the halfway level",
                 fontsize=14.5, fontweight="bold", pad=10)

    bed_y = 0.72
    x = np.linspace(0.65, 11.35, 500)
    surface = 1.38 + 2.55 / (1 + np.exp((x - 7.15) / 0.24))
    water_polygon = np.column_stack([np.r_[x, x[::-1]],
                                     np.r_[surface, np.full_like(x, bed_y)[::-1]]])
    ax.add_patch(Polygon(water_polygon, closed=True, facecolor=WATER, edgecolor="none"))
    ax.plot(x, surface, color=BLUE, linewidth=4)
    ax.plot([0.55, 11.45], [bed_y, bed_y], color=DARK, linewidth=4)

    midpoint_y = 2.655
    crossing_index = int(np.argmin(np.abs(surface - midpoint_y)))
    crossing_x = float(x[crossing_index])
    ax.plot([0.75, 11.15], [midpoint_y, midpoint_y], color=ORANGE,
            linewidth=3, linestyle="--")
    ax.scatter([crossing_x], [midpoint_y], marker="*", s=330, color=ORANGE,
               edgecolor="white", linewidth=1.3, zorder=6)
    ax.plot([crossing_x, crossing_x], [bed_y, midpoint_y], color=PURPLE,
            linewidth=3, linestyle=":")
    ax.add_patch(Polygon([[crossing_x, bed_y], [crossing_x + 0.22, bed_y - 0.25],
                          [crossing_x - 0.22, bed_y - 0.25]],
                         closed=True, facecolor=PURPLE, edgecolor="none"))

    ax.text(2.2, 4.18, "deeper water\nbehind the shock", ha="center",
            fontsize=12, color=BLUE, fontweight="bold")
    ax.text(9.85, 1.72, "shallow downstream\nwater", ha="center",
            fontsize=12, color=BLUE, fontweight="bold")
    ax.text(1.0, midpoint_y + 0.17, "halfway water level", fontsize=12,
            color=ORANGE, fontweight="bold")
    ax.annotate("the meeting point is the\npredicted shock location",
                xy=(crossing_x, midpoint_y), xytext=(9.2, 3.65),
                fontsize=12.5, ha="center", color=PURPLE, fontweight="bold",
                arrowprops=dict(arrowstyle="->", color=PURPLE, linewidth=2.2))
    ax.text(crossing_x, 0.12, "predicted shock\n3.7548 m", ha="center",
            va="top", fontsize=12, color=PURPLE, fontweight="bold")
    ax.annotate("flow direction", xy=(4.0, 1.05), xytext=(1.6, 1.05),
                va="center", fontsize=11, color=GRAY,
                arrowprops=dict(arrowstyle="->", color=GRAY, linewidth=1.8))

    fig.subplots_adjust(left=0.025, right=0.98, top=0.83, bottom=0.10)
    fig.savefig(FIGURE_B, dpi=240, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def make_error_illustration() -> None:
    fig, ax = plt.subplots(figsize=(7.2, 3.15))
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis("off")
    ax.set_title("(c) Compare the predicted shock with the correct shock location",
                 fontsize=15, fontweight="bold", pad=10)

    channel = FancyBboxPatch((0.65, 1.05), 10.7, 2.55,
                             boxstyle="round,pad=0.03,rounding_size=0.13",
                             facecolor=WATER, edgecolor=DARK, linewidth=2.5)
    ax.add_patch(channel)
    exact_x, predicted_x = 4.25, 7.55
    gap = Rectangle((exact_x, 1.08), predicted_x - exact_x, 2.49,
                    facecolor="#F6B8B8", edgecolor="none", alpha=0.72)
    gap.set_clip_path(channel)
    ax.add_patch(gap)
    ax.plot([exact_x, exact_x], [1.08, 3.57], color="#222222",
            linewidth=5, linestyle=":")
    ax.plot([predicted_x, predicted_x], [1.08, 3.57], color=PURPLE,
            linewidth=5, linestyle="-.")

    ax.annotate("correct shock\n3.7262 m", xy=(exact_x, 3.55), xytext=(2.75, 4.18),
                ha="center", fontsize=12, color="#222222", fontweight="bold",
                arrowprops=dict(arrowstyle="->", color="#222222", linewidth=1.8))
    ax.annotate("model's shock\n3.7548 m", xy=(predicted_x, 3.55), xytext=(9.15, 4.18),
                ha="center", fontsize=12, color=PURPLE, fontweight="bold",
                arrowprops=dict(arrowstyle="->", color=PURPLE, linewidth=1.8))
    ax.annotate("", xy=(exact_x, 2.34), xytext=(predicted_x, 2.34),
                arrowprops=dict(arrowstyle="<->", color=RED, linewidth=3.5))
    ax.text((exact_x + predicted_x) / 2, 2.62, "position error = 2.87 cm",
            ha="center", fontsize=15, color=RED, fontweight="bold")
    ax.text((exact_x + predicted_x) / 2, 1.68,
            "the red strip is how far the predicted front is displaced",
            ha="center", fontsize=11.5, color=DARK)
    ax.annotate("flow direction", xy=(10.55, 0.58), xytext=(8.4, 0.58),
                va="center", fontsize=11, color=GRAY,
                arrowprops=dict(arrowstyle="->", color=GRAY, linewidth=1.8))
    fig.subplots_adjust(left=0.025, right=0.98, top=0.82, bottom=0.05)
    fig.savefig(FIGURE_C, dpi=240, facecolor="white", bbox_inches="tight")
    plt.close(fig)


def make_figures() -> None:
    FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    make_midpoint_illustration()
    make_crossing_illustration()
    make_error_illustration()


def replace_image(document, caption_prefix: str, image_path: Path, width_inches: float) -> None:
    caption = next(
        (p for p in document.paragraphs if p.style.name == "Image Caption" and p.text.startswith(caption_prefix)),
        None,
    )
    if caption is None:
        raise RuntimeError(f"Could not find caption {caption_prefix}")
    figure_element = caption._p.getprevious()
    blip = figure_element.xpath(".//a:blip")[0]
    relationship_id = blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )
    document.part.related_parts[relationship_id]._blob = image_path.read_bytes()
    shape = next(
        item
        for item in document.inline_shapes
        if item._inline.getparent().getparent().getparent() is figure_element
    )
    with Image.open(image_path) as image:
        ratio = image.height / image.width
    shape.width = Inches(width_inches)
    shape.height = Inches(width_inches * ratio)


def update_explanation(document) -> None:
    paragraph = next(
        (p for p in document.paragraphs if "Figures 3(a)-3(c) show the shock measurement" in p.text),
        None,
    )
    if paragraph is None:
        raise RuntimeError("Could not find the shock-measurement explanation")
    equations = paragraph._p.xpath("./m:oMath")
    paragraph.clear()
    paragraph.add_run("Final depth and velocity errors are reported separately because ")
    paragraph._p.append(deepcopy(equations[0]))
    paragraph.add_run(
        " can hide which field controls a difference. Figures 3(a)-3(c) explain the "
        "shock measurement with three physical illustrations. The first illustration "
        "uses water columns to identify the halfway depth. The second shows the predicted "
        "water surface inside the channel; its meeting point with the halfway level marks "
        "the predicted shock. The third shows the channel from above. The gap between the "
        "correct and predicted fronts is the position error, which is 2.87 cm for the "
        "displayed Roe-R-RAR seed 0 run. A result is accepted only when the predicted water "
        "surface passes downward through the halfway level once."
    )


def update_captions(document) -> None:
    captions = {
        "Figure 3(a).": (
            "Figure 3(a). Water-column illustration of the midpoint depth. The orange "
            "surface is placed halfway between the depth behind the shock and the shallow "
            "downstream depth."
        ),
        "Figure 3(b).": (
            "Figure 3(b). Channel illustration of shock detection. The predicted shock is "
            "where the blue predicted water surface meets the orange halfway level."
        ),
        "Figure 3(c).": (
            "Figure 3(c). Top-down channel illustration of position error. The shaded red "
            "gap separates the correct shock front from the model's predicted front."
        ),
    }
    for prefix, replacement in captions.items():
        paragraph = next(
            p for p in document.paragraphs
            if p.style.name == "Image Caption" and p.text.startswith(prefix)
        )
        paragraph.text = replacement


def revise_document() -> None:
    document = Document(SOURCE)
    update_explanation(document)
    replace_image(document, "Figure 3(a).", FIGURE_A, 5.95)
    replace_image(document, "Figure 3(b).", FIGURE_B, 5.95)
    replace_image(document, "Figure 3(c).", FIGURE_C, 5.95)
    update_captions(document)
    document.save(OUTPUT)


def main() -> None:
    make_figures()
    revise_document()
    print(f"Created {OUTPUT}")
    print(f"Created {FIGURE_A}")
    print(f"Created {FIGURE_B}")
    print(f"Created {FIGURE_C}")


if __name__ == "__main__":
    main()
